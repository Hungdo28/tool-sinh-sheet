# -*- coding: utf-8 -*-
"""
Excel Processor Module
Xử lý đọc file Test Case tổng, phân tích dữ liệu, nhóm test case con và sinh ra các sheet chi tiết.
"""

import os
import re
import urllib.request
import urllib.error
from copy import copy
from typing import Dict, List, Any, Tuple, Optional, Callable
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def download_google_sheet_as_excel(url: str, save_path: str) -> str:
    """
    Tải file Excel từ link Google Sheets (yêu cầu Sheet ở chế độ Bất kỳ ai có liên kết đều có thể xem).
    """
    match = re.search(r'/spreadsheets/d/([a-zA-Z0-9-_]+)', url)
    if not match:
        raise ValueError("Đường dẫn Google Sheets không hợp lệ! Định dạng chuẩn: https://docs.google.com/spreadsheets/d/.../edit")

    sheet_id = match.group(1)
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"

    req = urllib.request.Request(
        export_url,
        headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
    )

    try:
        os.makedirs(os.path.dirname(os.path.abspath(save_path)), exist_ok=True)
        with urllib.request.urlopen(req, timeout=15) as response:
            content = response.read()
            # Kiểm tra xem có trả về file zip/excel (PK..) hay trang HTML đăng nhập
            if not content.startswith(b'PK'):
                raise ValueError("Không thể tải file! Vui lòng đảm bảo Google Sheet đã được 'Chia sẻ' ở chế độ 'Bất kỳ ai có liên kết đều có thể xem'.")
            with open(save_path, 'wb') as out_f:
                out_f.write(content)
        return save_path
    except urllib.error.HTTPError as e:
        if e.code == 404:
            raise ValueError("Không tìm thấy Google Sheet. Vui lòng kiểm tra lại link.")
        elif e.code in (401, 403):
            raise ValueError("Google Sheet này chưa mở quyền truy cập. Vui lòng bật 'Bất kỳ ai có liên kết' trong menu Chia sẻ.")
        else:
            raise ValueError(f"Lỗi kết nối tới Google Sheets: HTTP {e.code}")
    except Exception as e:
        raise ValueError(f"Lỗi khi tải Google Sheet: {str(e)}")



def sanitize_sheet_name(name: Any, existing_names: Optional[List[str]] = None) -> str:
    """
    Chuẩn hóa tên Sheet theo tiêu chuẩn Excel:
    - Loại bỏ ký tự cấm: \\ / ? * [ ] :
    - Tối đa 31 ký tự
    - Không bắt đầu/kết thúc bằng dấu nháy đơn '
    - Không trùng lặp (không phân biệt hoa thường)
    """
    if name is None or str(name).strip() == "":
        base_name = "Chua_Phan_Nhom"
    else:
        # Chuyển thành chuỗi và loại bỏ ký tự cấm của Excel
        base_name = re.sub(r'[\\/*?:\[\]]', '_', str(name).strip())
        # Loại bỏ dấu nháy đơn ở đầu và cuối
        base_name = base_name.strip("'").strip()
        if not base_name:
            base_name = "Sheet"

    # Giới hạn tối đa 31 ký tự
    clean_name = base_name[:31]

    if existing_names is None:
        return clean_name

    # Đảm bảo tính duy nhất (Excel không phân biệt chữ hoa/thường)
    existing_lower = {n.lower() for n in existing_names}
    if clean_name.lower() not in existing_lower:
        return clean_name

    # Nếu trùng, thêm hậu tố _1, _2...
    counter = 1
    while True:
        suffix = f"_{counter}"
        avail_len = 31 - len(suffix)
        candidate = f"{clean_name[:avail_len]}{suffix}"
        if candidate.lower() not in existing_lower:
            return candidate
        counter += 1


def get_cell_value_with_merged_resolution(sheet, row: int, col: int) -> Any:
    """
    Lấy giá trị của ô, nếu ô nằm trong vùng Merged Cells thì lấy giá trị của ô gốc (top-left).
    """
    cell = sheet.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value

    # Kiểm tra ô có nằm trong dải merge nào không
    coord = cell.coordinate
    for merged_range in sheet.merged_cells.ranges:
        if coord in merged_range:
            top_left_cell = sheet.cell(row=merged_range.min_row, column=merged_range.min_col)
            return top_left_cell.value

    return None


def inspect_excel_file(file_path: str) -> Dict[str, Any]:
    """
    Kiểm tra nhanh file Excel: danh sách sheet, sheet mặc định.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    active_sheet = sheet_names[0] if sheet_names else ""
    wb.close()

    return {
        "sheet_names": sheet_names,
        "active_sheet": active_sheet
    }


def get_sheet_columns_preview(file_path: str, sheet_name: str, header_row: int = 1) -> Dict[str, Any]:
    """
    Đọc danh sách cột và các dòng mẫu để hiển thị cho người dùng chọn cột phân nhóm.
    Hỗ trợ kết hợp Header đa tầng (ví dụ: Data test - Lần 1).
    """
    wb = openpyxl.load_workbook(file_path, read_only=False, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' không tồn tại trong file.")

    ws = wb[sheet_name]
    max_col = ws.max_column
    max_row = ws.max_row

    # Xác định dòng chứa tên cột chính và dòng nhóm cha
    actual_col_row = header_row
    parent_row = 1 if header_row > 1 else None

    # Nếu header_row >= 3, kiểm tra xem header_row có phải là dòng banner (chỉ có 1-2 ô có chữ)
    if header_row >= 3:
        filled_in_h = sum(1 for c in range(1, max_col + 1) if ws.cell(row=header_row, column=c).value is not None)
        filled_in_prev = sum(1 for c in range(1, max_col + 1) if ws.cell(row=header_row - 1, column=c).value is not None)
        if filled_in_h <= 3 and filled_in_prev > filled_in_h:
            # Dòng header_row là dòng banner (ví dụ: Job quét nợ), tên cột thực tế nằm ở dòng header_row - 1
            actual_col_row = header_row - 1
            parent_row = 1

    columns = []
    for c in range(1, max_col + 1):
        col_name = get_cell_value_with_merged_resolution(ws, actual_col_row, c)
        
        # Kiểm tra xem dòng cha có tên nhóm (như Data test, UAT, PROD)
        parent_header = ""
        if parent_row and parent_row != actual_col_row:
            p_val = get_cell_value_with_merged_resolution(ws, parent_row, c)
            if p_val and str(p_val).strip() != str(col_name).strip():
                parent_header = str(p_val).strip()

        if col_name is None or str(col_name).strip() == "":
            if parent_header:
                full_name = parent_header
            else:
                full_name = f"Cột {get_column_letter(c)}"
        else:
            if parent_header and str(col_name).strip() not in parent_header:
                full_name = f"{parent_header} - {str(col_name).strip()}"
            else:
                full_name = str(col_name).strip()

        columns.append({
            "index": c,
            "letter": get_column_letter(c),
            "name": full_name
        })

    # Lấy các dòng xem trước (tối đa 150 dòng)
    preview_rows = []
    start_row = header_row + 1
    end_row = min(start_row + 150, max_row + 1)
    for r in range(start_row, end_row):
        row_vals = []
        has_content = False
        for c in range(1, max_col + 1):
            val = get_cell_value_with_merged_resolution(ws, r, c)
            val_str = "" if val is None else str(val).strip()
            if val_str:
                has_content = True
            row_vals.append(val_str)
        if has_content:
            preview_rows.append(row_vals)

    wb.close()
    return {
        "columns": columns,
        "total_rows": max_row,
        "data_rows_count": max(0, max_row - header_row),
        "preview_rows": preview_rows
    }


def analyze_groups(
    file_path: str,
    sheet_name: str,
    group_column_idx: int,
    header_row: int = 1,
    forward_fill: bool = True
) -> Dict[str, Any]:
    """
    Phân tích các nhóm testcase dựa trên cột được chọn:
    - Tìm các giá trị nhóm duy nhất
    - Đếm số lượng testcase con trong từng nhóm
    - Chuẩn hóa tên sheet tương ứng
    - Hỗ trợ forward fill cho các ô gộp hoặc để trống
    """
    wb = openpyxl.load_workbook(file_path, read_only=False, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' không tồn tại trong file.")

    ws = wb[sheet_name]
    max_row = ws.max_row

    groups: Dict[str, Dict[str, Any]] = {}
    last_group_name = None
    existing_sheet_names: List[str] = ["Tổng quan", "Overview", "Index"]

    # Quét qua các dòng dữ liệu
    for r in range(header_row + 1, max_row + 1):
        # Kiểm tra xem dòng có rỗng hoàn toàn không
        has_data = False
        for c in range(1, min(ws.max_column + 1, 15)):
            if ws.cell(row=r, column=c).value is not None:
                has_data = True
                break
        if not has_data:
            continue

        raw_group_val = get_cell_value_with_merged_resolution(ws, r, group_column_idx)
        val_clean = str(raw_group_val).strip() if raw_group_val is not None else ""

        # Bỏ qua nếu giá trị trùng với tiêu đề cột lặp lại
        if val_clean.lower() in ("id", "mã test case", "stt", "mã tc", "testcase id", "tên test case", "mục đích kiểm thử"):
            continue

        if val_clean != "":
            group_name = val_clean
            last_group_name = group_name
        elif forward_fill and last_group_name is not None:
            group_name = last_group_name
        else:
            # Bỏ qua các dòng banner phụ trước khi có testcase ID đầu tiên
            continue

        if group_name not in groups:
            sheet_name_clean = sanitize_sheet_name(group_name, existing_sheet_names)
            existing_sheet_names.append(sheet_name_clean)
            groups[group_name] = {
                "group_name": group_name,
                "sheet_name": sheet_name_clean,
                "count": 0,
                "row_indices": []
            }

        groups[group_name]["count"] += 1
        groups[group_name]["row_indices"].append(r)

    wb.close()
    return groups


def fix_fill_alpha(fill):
    """
    Chuẩn hóa Alpha channel của màu ô thành 'FF' (100% Opacity đặc).
    Tránh lỗi openpyxl tự động thêm '00' khiến Excel hiển thị ô bị trong suốt/mất màu.
    """
    if not fill or not fill.fill_type:
        return fill
    f = copy(fill)
    try:
        if hasattr(f, 'start_color') and f.start_color and f.start_color.type == 'rgb':
            rgb = str(f.start_color.rgb)
            if len(rgb) == 8 and rgb.startswith('00') and rgb != '00000000':
                f.start_color.rgb = 'FF' + rgb[2:]
            elif len(rgb) == 6:
                f.start_color.rgb = 'FF' + rgb
        if hasattr(f, 'end_color') and f.end_color and f.end_color.type == 'rgb':
            rgb = str(f.end_color.rgb)
            if len(rgb) == 8 and rgb.startswith('00') and rgb != '00000000':
                f.end_color.rgb = 'FF' + rgb[2:]
            elif len(rgb) == 6:
                f.end_color.rgb = 'FF' + rgb
    except Exception:
        pass
    return f


def copy_cell_style_and_value(src_cell, dst_cell, copy_style: bool = True):
    """
    Sao chép giá trị và định dạng từ ô nguồn sang ô đích.
    Tái tạo trực tiếp PatternFill, Font, Border, Alignment độc lập để đảm bảo
    không bị lỗi StyleProxy và bắt buộc Alpha = FF (100% Opaque) để Excel luôn đổ màu rực rỡ.
    """
    dst_cell.value = src_cell.value
    if src_cell.comment:
        dst_cell.comment = copy(src_cell.comment)

    if not copy_style or not src_cell.has_style:
        return

    # 1. Fill / Màu nền ô (Bắt buộc Alpha = FF để không bị Excel làm trong suốt)
    if src_cell.fill and src_cell.fill.fill_type:
        try:
            f = src_cell.fill
            start_c = str(f.start_color.rgb) if f.start_color and f.start_color.type == 'rgb' else None
            end_c = str(f.end_color.rgb) if f.end_color and f.end_color.type == 'rgb' else start_c

            if start_c:
                if len(start_c) == 8 and start_c.startswith('00') and start_c != '00000000':
                    start_c = 'FF' + start_c[2:]
                elif len(start_c) == 6:
                    start_c = 'FF' + start_c
            if end_c:
                if len(end_c) == 8 and end_c.startswith('00') and end_c != '00000000':
                    end_c = 'FF' + end_c[2:]
                elif len(end_c) == 6:
                    end_c = 'FF' + end_c

            if start_c:
                dst_cell.fill = PatternFill(fill_type=f.fill_type, start_color=start_c, end_color=end_c)
            elif f.fill_type:
                dst_cell.fill = PatternFill(fill_type=f.fill_type)
        except Exception:
            pass

    # 2. Font / Kiểu chữ
    if src_cell.font:
        try:
            f = src_cell.font
            f_color = f.color
            if f_color and f_color.type == 'rgb':
                c = str(f_color.rgb)
                if len(c) == 8 and c.startswith('00') and c != '00000000':
                    f_color = 'FF' + c[2:]
                elif len(c) == 6:
                    f_color = 'FF' + c
            dst_cell.font = Font(
                name=f.name,
                size=f.size,
                bold=f.bold,
                italic=f.italic,
                underline=f.underline,
                strike=f.strike,
                color=f_color
            )
        except Exception:
            pass

    # 3. Alignment / Căn lề
    if src_cell.alignment:
        try:
            a = src_cell.alignment
            dst_cell.alignment = Alignment(
                horizontal=a.horizontal,
                vertical=a.vertical,
                wrap_text=a.wrap_text,
                shrink_to_fit=a.shrink_to_fit,
                indent=a.indent
            )
        except Exception:
            pass

    # 4. Border / Viền kẻ bảng
    if src_cell.border:
        try:
            b = src_cell.border
            def _copy_side(s):
                if not s or not s.style:
                    return None
                c = s.color.rgb if s.color and s.color.type == 'rgb' else 'FF000000'
                if len(str(c)) == 8 and str(c).startswith('00') and str(c) != '00000000':
                    c = 'FF' + str(c)[2:]
                elif len(str(c)) == 6:
                    c = 'FF' + str(c)
                return Side(style=s.style, color=c)

            dst_cell.border = Border(
                left=_copy_side(b.left),
                right=_copy_side(b.right),
                top=_copy_side(b.top),
                bottom=_copy_side(b.bottom)
            )
        except Exception:
            pass

    # 5. Number format
    if src_cell.number_format:
        dst_cell.number_format = src_cell.number_format


def generate_excel_sheets(
    input_file: str,
    output_file: str,
    source_sheet_name: str,
    group_column_idx: int,
    header_row: int = 1,
    selected_group_keys: Optional[List[str]] = None,
    forward_fill: bool = True,
    copy_styles: bool = True,
    include_summary_sheet: bool = True,
    add_back_link: bool = True,
    export_mode: str = "single_workbook",
    progress_callback: Optional[Callable[[float, str], None]] = None
) -> Dict[str, Any]:
    """
    Hàm chính để sinh ra file Excel với từng sheet theo nhóm testcase con.
    """
    if progress_callback:
        progress_callback(5.0, "Đang mở file Excel nguồn...")

    # Load workbook nguồn
    src_wb = openpyxl.load_workbook(input_file, data_only=False)
    if source_sheet_name not in src_wb.sheetnames:
        raise ValueError(f"Sheet '{source_sheet_name}' không tìm thấy trong file nguồn.")
    src_ws = src_wb[source_sheet_name]

    max_col = src_ws.max_column
    max_row = src_ws.max_row

    if progress_callback:
        progress_callback(15.0, "Đang phân tích các nhóm test case...")

    # Phân tích danh sách nhóm
    groups = analyze_groups(
        input_file,
        source_sheet_name,
        group_column_idx,
        header_row,
        forward_fill
    )

    if not groups:
        raise ValueError("Không tìm thấy dòng dữ liệu testcase nào để phân nhóm.")

    # Lọc các nhóm được chọn
    if selected_group_keys:
        groups = {k: v for k, v in groups.items() if k in selected_group_keys}

    if not groups:
        raise ValueError("Bạn chưa chọn nhóm testcase nào để xuất sheet.")

    total_groups = len(groups)
    created_sheets_info = []

    if export_mode == "single_workbook":
        dst_wb = openpyxl.Workbook()
        # Xóa sheet mặc định ban đầu
        default_sheet = dst_wb.active

        summary_sheet_name = "Tổng quan"
        summary_ws = None
        if include_summary_sheet:
            summary_ws = dst_wb.create_sheet(title=summary_sheet_name)

        # Định dạng style cho bảng tổng quan (Dùng FF prefix để màu hiển thị đặc 100% trên Excel)
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFFFF")
        header_fill = PatternFill(start_color="FF1E3A8A", end_color="FF1E3A8A", fill_type="solid")  # Navy Blue
        title_font = Font(name="Segoe UI", size=16, bold=True, color="FF1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="FF475569")
        data_font = Font(name="Segoe UI", size=10)
        link_font = Font(name="Segoe UI", size=10, color="FF2563EB", underline="single")
        alt_fill = PatternFill(start_color="FFF8FAFC", end_color="FFF8FAFC", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='FFCBD5E1'),
            right=Side(style='thin', color='FFCBD5E1'),
            top=Side(style='thin', color='FFCBD5E1'),
            bottom=Side(style='thin', color='FFCBD5E1')
        )

        # Tạo từng sheet cho mỗi nhóm
        group_items = list(groups.items())
        for idx, (group_name, group_data) in enumerate(group_items):
            current_progress = 20.0 + (idx / total_groups) * 60.0
            sheet_title = group_data["sheet_name"]
            if progress_callback:
                progress_callback(
                    current_progress,
                    f"Đang tạo sheet [{idx + 1}/{total_groups}]: {sheet_title} ({group_data['count']} testcase)..."
                )

            new_ws = dst_wb.create_sheet(title=sheet_title)
            created_sheets_info.append({
                "group_name": group_name,
                "sheet_name": sheet_title,
                "count": group_data["count"]
            })

            # Sao chép độ rộng cột (column dimensions)
            for c in range(1, max_col + 1):
                col_letter = get_column_letter(c)
                if col_letter in src_ws.column_dimensions:
                    new_ws.column_dimensions[col_letter].width = src_ws.column_dimensions[col_letter].width

            # Header template bắt đầu ngay từ dòng 1 (Khớp 100% template gốc, không chèn dòng quay lại)
            target_row = 1
            header_target_start = 1
            # Sao chép các dòng Header từ nguồn
            for r in range(1, header_row + 1):
                if r in src_ws.row_dimensions and src_ws.row_dimensions[r].height:
                    new_ws.row_dimensions[target_row].height = src_ws.row_dimensions[r].height
                for c in range(1, max_col + 1):
                    src_cell = src_ws.cell(row=r, column=c)
                    dst_cell = new_ws.cell(row=target_row, column=c)
                    copy_cell_style_and_value(src_cell, dst_cell, copy_style=copy_styles)
                target_row += 1

            # Tái tạo các ô gộp (Merged Cells) thuộc phần Header template và đồng bộ màu ô
            offset = header_target_start - 1
            for rng in src_ws.merged_cells.ranges:
                if rng.max_row <= header_row:
                    m_min_r = rng.min_row + offset
                    m_max_r = rng.max_row + offset
                    m_min_c = rng.min_col
                    m_max_c = rng.max_col
                    new_ws.merge_cells(
                        start_row=m_min_r,
                        start_column=m_min_c,
                        end_row=m_max_r,
                        end_column=m_max_c
                    )
                    top_cell = new_ws.cell(row=m_min_r, column=m_min_c)
                    if copy_styles and top_cell.has_style:
                        for mr in range(m_min_r, m_max_r + 1):
                            for mc in range(m_min_c, m_max_c + 1):
                                if mr == m_min_r and mc == m_min_c:
                                    continue
                                c_cell = new_ws.cell(row=mr, column=mc)
                                if top_cell.fill and top_cell.fill.fill_type:
                                    c_cell.fill = copy(top_cell.fill)
                                if top_cell.border:
                                    c_cell.border = copy(top_cell.border)

            # Sao chép các dòng dữ liệu của nhóm
            for src_row_idx in group_data["row_indices"]:
                if src_row_idx in src_ws.row_dimensions and src_ws.row_dimensions[src_row_idx].height:
                    new_ws.row_dimensions[target_row].height = src_ws.row_dimensions[src_row_idx].height

                for c in range(1, max_col + 1):
                    src_cell = src_ws.cell(row=src_row_idx, column=c)
                    dst_cell = new_ws.cell(row=target_row, column=c)
                    copy_cell_style_and_value(src_cell, dst_cell, copy_style=copy_styles)

                    # Nếu ô nhóm ở dòng này bị None do merge ở file gốc, điền lại tên nhóm rõ ràng
                    if c == group_column_idx and (dst_cell.value is None or str(dst_cell.value).strip() == ""):
                        dst_cell.value = group_name

                target_row += 1

        # Xây dựng trang Tổng quan (Summary / Overview Sheet)
        if include_summary_sheet and summary_ws:
            if progress_callback:
                progress_callback(85.0, "Đang thiết lập bảng tổng quan và liên kết...")

            # Tiêu đề trang Tổng quan
            summary_ws.cell(row=2, column=2, value="BẢNG TỔNG QUAN TEST CASES THEO MODULE / NHÓM").font = title_font
            summary_ws.cell(
                row=3, column=2,
                value=f"File nguồn: {os.path.basename(input_file)} | Tổng số nhóm: {total_groups} | Tổng số testcase: {sum(g['count'] for g in groups.values())}"
            ).font = sub_font

            # Header bảng tổng quan
            headers = ["STT", "Tên Phân Nhóm / Module", "Tên Sheet Chi Tiết (Nhấp để mở)", "Số Test Case Con", "Tỷ Lệ (%)"]
            header_row_idx = 5
            summary_ws.row_dimensions[header_row_idx].height = 28

            for col_i, h_name in enumerate(headers, start=2):
                h_cell = summary_ws.cell(row=header_row_idx, column=col_i, value=h_name)
                h_cell.font = header_font
                h_cell.fill = header_fill
                h_cell.alignment = center_align
                h_cell.border = thin_border

            total_tc_count = sum(g["count"] for g in groups.values())

            # Dòng dữ liệu tổng quan
            curr_row = header_row_idx + 1
            for stt, (g_name, g_info) in enumerate(groups.items(), start=1):
                summary_ws.row_dimensions[curr_row].height = 22
                fill_to_use = alt_fill if stt % 2 == 0 else None

                # STT
                c_stt = summary_ws.cell(row=curr_row, column=2, value=stt)
                c_stt.alignment = center_align
                c_stt.font = data_font
                c_stt.border = thin_border

                # Tên nhóm
                c_name = summary_ws.cell(row=curr_row, column=3, value=g_name)
                c_name.alignment = left_align
                c_name.font = data_font
                c_name.border = thin_border

                # Link đến Sheet
                sheet_target = g_info["sheet_name"]
                c_link = summary_ws.cell(row=curr_row, column=4, value=f"📑 {sheet_target}")
                c_link.hyperlink = f"#'{sheet_target}'!A1"
                c_link.alignment = left_align
                c_link.font = link_font
                c_link.border = thin_border

                # Số lượng test case
                c_count = summary_ws.cell(row=curr_row, column=5, value=g_info["count"])
                c_count.alignment = center_align
                c_count.font = data_font
                c_count.border = thin_border

                # Tỷ lệ %
                pct = (g_info["count"] / total_tc_count * 100) if total_tc_count > 0 else 0
                c_pct = summary_ws.cell(row=curr_row, column=6, value=f"{pct:.1f}%")
                c_pct.alignment = center_align
                c_pct.font = data_font
                c_pct.border = thin_border

                if fill_to_use:
                    for c_idx in range(2, 7):
                        summary_ws.cell(row=curr_row, column=c_idx).fill = fill_to_use

                curr_row += 1

            # Dòng tổng cộng
            summary_ws.row_dimensions[curr_row].height = 24
            c_sum_label = summary_ws.cell(row=curr_row, column=2, value="TỔNG CỘNG")
            c_sum_label.font = Font(name="Segoe UI", size=10, bold=True)
            c_sum_label.alignment = center_align
            c_sum_label.border = thin_border

            summary_ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=4)

            c_sum_val = summary_ws.cell(row=curr_row, column=5, value=total_tc_count)
            c_sum_val.font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
            c_sum_val.alignment = center_align
            c_sum_val.border = thin_border

            c_sum_pct = summary_ws.cell(row=curr_row, column=6, value="100.0%")
            c_sum_pct.font = Font(name="Segoe UI", size=10, bold=True)
            c_sum_pct.alignment = center_align
            c_sum_pct.border = thin_border

            # Tự căn chỉnh chiều rộng cột trang Tổng quan
            summary_ws.column_dimensions['A'].width = 3
            summary_ws.column_dimensions['B'].width = 8
            summary_ws.column_dimensions['C'].width = 38
            summary_ws.column_dimensions['D'].width = 35
            summary_ws.column_dimensions['E'].width = 20
            summary_ws.column_dimensions['F'].width = 16

        # Xóa sheet rỗng mặc định nếu còn tồn tại
        if default_sheet in dst_wb.worksheets:
            dst_wb.remove(default_sheet)

        if progress_callback:
            progress_callback(92.0, "Đang lưu file Excel kết quả...")

        # Tạo thư mục đích nếu chưa có
        os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
        dst_wb.save(output_file)
        dst_wb.close()

    elif export_mode == "multiple_files":
        # Chế độ xuất mỗi nhóm thành 1 file riêng
        output_dir = output_file
        os.makedirs(output_dir, exist_ok=True)

        group_items = list(groups.items())
        for idx, (group_name, group_data) in enumerate(group_items):
            current_progress = 20.0 + (idx / total_groups) * 70.0
            sheet_title = group_data["sheet_name"]
            single_file_name = os.path.join(output_dir, f"{sheet_title}.xlsx")

            if progress_callback:
                progress_callback(
                    current_progress,
                    f"Đang xuất file [{idx + 1}/{total_groups}]: {os.path.basename(single_file_name)}..."
                )

            single_wb = openpyxl.Workbook()
            single_ws = single_wb.active
            single_ws.title = sheet_title

            # Sao chép độ rộng cột
            for c in range(1, max_col + 1):
                col_letter = get_column_letter(c)
                if col_letter in src_ws.column_dimensions:
                    single_ws.column_dimensions[col_letter].width = src_ws.column_dimensions[col_letter].width

            target_row = 1
            header_target_start = target_row
            # Header
            for r in range(1, header_row + 1):
                if r in src_ws.row_dimensions and src_ws.row_dimensions[r].height:
                    single_ws.row_dimensions[target_row].height = src_ws.row_dimensions[r].height
                for c in range(1, max_col + 1):
                    src_cell = src_ws.cell(row=r, column=c)
                    dst_cell = single_ws.cell(row=target_row, column=c)
                    copy_cell_style_and_value(src_cell, dst_cell, copy_style=copy_styles)
                target_row += 1

            # Tái tạo các ô gộp thuộc phần Header template
            offset = header_target_start - 1
            for rng in src_ws.merged_cells.ranges:
                if rng.max_row <= header_row:
                    single_ws.merge_cells(
                        start_row=rng.min_row + offset,
                        start_column=rng.min_col,
                        end_row=rng.max_row + offset,
                        end_column=rng.max_col
                    )

            # Dữ liệu
            for src_row_idx in group_data["row_indices"]:
                if src_row_idx in src_ws.row_dimensions and src_ws.row_dimensions[src_row_idx].height:
                    single_ws.row_dimensions[target_row].height = src_ws.row_dimensions[src_row_idx].height
                for c in range(1, max_col + 1):
                    src_cell = src_ws.cell(row=src_row_idx, column=c)
                    dst_cell = single_ws.cell(row=target_row, column=c)
                    copy_cell_style_and_value(src_cell, dst_cell, copy_style=copy_styles)
                    if c == group_column_idx and (dst_cell.value is None or str(dst_cell.value).strip() == ""):
                        dst_cell.value = group_name
                target_row += 1

            single_wb.save(single_file_name)
            single_wb.close()

            created_sheets_info.append({
                "group_name": group_name,
                "sheet_name": sheet_title,
                "file_path": single_file_name,
                "count": group_data["count"]
            })

    src_wb.close()

    if progress_callback:
        progress_callback(100.0, "Hoàn tất sinh sheet thành công!")

    return {
        "status": "success",
        "total_groups": total_groups,
        "total_testcases": sum(g["count"] for g in groups.values()),
        "output_path": output_file,
        "sheets": created_sheets_info
    }
